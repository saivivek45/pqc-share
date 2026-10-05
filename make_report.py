"""Fill 'Project_Report_Template .docx' with the PQC-Share project report."""
import copy
from pathlib import Path

import docx
from docx.shared import Pt

ROOT = Path(__file__).parent
TITLE = "PQC-Share: Post-Quantum Secure File Sharing System"

d = docx.Document(ROOT / "Project_Report_Template .docx")
P = d.paragraphs


def set_text(p, text):
    """Replace paragraph text keeping the first run's formatting."""
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


def add_after(p, text="", style="Normal", code=False, bold=False):
    new = copy.deepcopy(p._p)
    p._p.addnext(new)
    np_ = docx.text.paragraph.Paragraph(new, p._parent)
    for r in list(np_.runs):
        r._r.getparent().remove(r._r)
    np_.style = d.styles[style]
    run = np_.add_run(text)
    run.bold = bold
    if code:
        run.font.name = "Courier New"
        run.font.size = Pt(8)
    return np_


def fill(p, items):
    """items: list of str or ('code'|'bold', str). First goes into p."""
    cur = p
    first = True
    for it in items:
        kind, text = ("text", it) if isinstance(it, str) else it
        if first:
            set_text(cur, text)
            if kind == "code":
                cur.runs[0].font.name = "Courier New"
                cur.runs[0].font.size = Pt(8)
            cur.runs[0].bold = kind == "bold"
            first = False
        else:
            cur = add_after(cur, text, code=kind == "code", bold=kind == "bold")


def read(rel):
    return (ROOT / rel).read_text(encoding="utf-8").rstrip()


# ---- cover page ----
set_text(P[4], TITLE)
set_text(P[5], "Course Name: Cryptography / Information Security")
set_text(P[6], "Student Name: Vivek")

# ---- body (grab refs first; inserting shifts nothing in the P list) ----
set_text(P[20], TITLE)

fill(P[21], [
    "A full-stack web application that lets users exchange files securely using "
    "NIST-standardised post-quantum cryptography (ML-KEM-768 and ML-DSA-65) combined "
    "with AES-256-GCM."
])

fill(P[24], [
    "Large-scale quantum computers running Shor's algorithm would break RSA and "
    "elliptic-curve cryptography, and adversaries can already record encrypted traffic "
    "today to decrypt it later (\"harvest now, decrypt later\"). PQC-Share is a secure "
    "file-sharing system that is resistant to such attacks. Files are encrypted with "
    "AES-256-GCM using a session key established through the ML-KEM-768 key "
    "encapsulation mechanism (NIST FIPS 203), and every package is authenticated with "
    "an ML-DSA-65 digital signature (NIST FIPS 204). The backend is built with FastAPI "
    "and SQLAlchemy, the frontend with React and TypeScript. The system also includes "
    "an audit log, benchmarking of the post-quantum scheme against an RSA baseline, and "
    "a test suite covering the cryptographic core and the API."
])

fill(P[27], [
    "Public-key schemes in use today (RSA, ECDH, ECDSA) rely on integer factoring and "
    "discrete logarithms, both solvable in polynomial time by a sufficiently large "
    "quantum computer. NIST finalised the first post-quantum standards in 2024: "
    "ML-KEM (FIPS 203, formerly CRYSTALS-Kyber) for key establishment and ML-DSA "
    "(FIPS 204, formerly CRYSTALS-Dilithium) for signatures.",
    "Objectives of this project:",
    "1. Build a file-sharing application where sender and receiver exchange files "
    "protected by post-quantum primitives.",
    "2. Provide confidentiality (AES-256-GCM), integrity and sender authenticity "
    "(ML-DSA-65 signature plus the GCM tag).",
    "3. Protect users' private keys at rest by wrapping them with a master key.",
    "4. Compare the performance and key/ciphertext sizes of the PQC scheme with a "
    "classical RSA baseline.",
    "Technology stack: Python 3.12, FastAPI, SQLAlchemy 2, SQLite, pqcrypto (PQClean), "
    "cryptography, Argon2 password hashing, PyJWT; React 19, TypeScript, Vite, "
    "Tailwind CSS, Recharts; pytest and Vitest for testing.",
])

fill(P[30], [
    ("bold", "1. Key management"),
    "On registration each user receives an ML-KEM-768 key pair (encryption) and an "
    "ML-DSA-65 key pair (signing). Passwords are hashed with Argon2. Private keys are "
    "never stored in the clear: each is encrypted with AES-256-GCM under a key derived "
    "via HKDF-SHA256 from a server master key, bound to the user id and the key "
    "fingerprint (SHA3-256 of the public key).",
    ("bold", "2. Hybrid encryption (sender)"),
    "a) Encapsulate against the receiver's ML-KEM public key to obtain a KEM ciphertext "
    "and a shared secret.  b) Derive a 256-bit session key with HKDF-SHA256, bound to the "
    "package id and both fingerprints.  c) Encrypt the file with AES-256-GCM (random "
    "96-bit nonce); the canonical JSON header is used as additional authenticated data.  "
    "d) Sign DOMAIN || header || SHA3-256(ciphertext) with the sender's ML-DSA-65 key.",
    ("bold", "3. Package format"),
    "Container: magic \"PQSH\" | u32 header length | JSON header | ciphertext || GCM tag. "
    "The header carries the version, package id, filename, size, timestamp, algorithm "
    "identifiers, nonce, KEM ciphertext, sender/receiver fingerprints and the signature. "
    "Any modification of a header byte or the body invalidates both the signature and "
    "the GCM tag.",
    ("bold", "4. Decryption (receiver)"),
    "Parse and validate the package, verify the sender's signature first, check that the "
    "package was addressed to the receiver's key, decapsulate with the ML-KEM private "
    "key, re-derive the session key and perform authenticated AES-GCM decryption. "
    "Plaintext is returned only if every check passes.",
    ("bold", "5. System architecture"),
    "React single-page app -> REST API (FastAPI, JWT authentication) -> services "
    "(keys, file service, audit, benchmark) -> crypto layer (kem.py, sig.py, "
    "primitives.py, package.py) and SQLite/file storage. The KEM and signature "
    "algorithms sit behind abstract backends so they can be swapped (e.g. liboqs, "
    "ML-KEM-1024).",
    ("bold", "6. Evaluation"),
    "The benchmark module times key generation, encapsulation/decapsulation, "
    "sign/verify and AES operations, and compares them with RSA (rsa_baseline.py). "
    "Results are shown as charts on the Benchmarks and Comparison pages. Automated "
    "tests in backend/tests check round-trip correctness, tamper detection, wrong-receiver "
    "rejection and the API endpoints.",
])

fill(P[33], [
    "The web application consists of the following screens. Insert a screenshot of each "
    "below its caption.",
    ("bold", "Figure 1: Login / Registration page (AuthPage)"),
    ("bold", "Figure 2: Dashboard with key fingerprints and activity (Dashboard)"),
    ("bold", "Figure 3: Encrypting and sharing a file with another user (Share)"),
    ("bold", "Figure 4: Incoming/outgoing transfers, verify and decrypt (Transfers)"),
    ("bold", "Figure 5: Performance benchmarks of ML-KEM / ML-DSA / AES (Benchmarks)"),
    ("bold", "Figure 6: PQC versus RSA comparison (Comparison)"),
])

code_files = [
    ("backend/app/crypto/primitives.py", "Symmetric layer: AES-256-GCM, HKDF, fingerprints, key wrapping"),
    ("backend/app/crypto/kem.py", "ML-KEM-768 key encapsulation backend"),
    ("backend/app/crypto/sig.py", "ML-DSA-65 signature backend"),
    ("backend/app/crypto/package.py", "Package format and encrypt / verify / decrypt workflow"),
]
items = ["Core cryptographic modules (full source in the project repository):"]
for rel, desc in code_files:
    items.append(("bold", f"{rel} - {desc}"))
    items.append(("code", read(rel)))
fill(P[36], items)

fill(P[38], [
    "[1] NIST, FIPS 203: Module-Lattice-Based Key-Encapsulation Mechanism Standard, 2024.",
    "[2] NIST, FIPS 204: Module-Lattice-Based Digital Signature Standard, 2024.",
    "[3] NIST, SP 800-38D: Recommendation for GCM and GMAC, 2007.",
    "[4] H. Krawczyk and P. Eronen, RFC 5869: HKDF, IETF, 2010.",
    "[5] P. Shor, Polynomial-Time Algorithms for Prime Factorization and Discrete "
    "Logarithms on a Quantum Computer, SIAM J. Computing, 1997.",
    "[6] PQClean project: https://github.com/PQClean/PQClean",
    "[7] FastAPI documentation: https://fastapi.tiangolo.com",
])

out = ROOT / "PQC_Share_Project_Report.docx"
d.save(out)
print("saved", out)
